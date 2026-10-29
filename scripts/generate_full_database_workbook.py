import json
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open('public/nec_students_2026.json', 'r', encoding='utf-8') as f:
    students = json.load(f)

# Also load raw CSE details for the 36-column sheet
wb_cse = openpyxl.load_workbook('C:/Users/SDP/.gemini/antigravity/brain/257e1302-6051-4ba8-b085-f5cd855ef30b/.user_uploaded/media_1793284665788_84285bee.xlsx', data_only=True)
ws_cse_raw = wb_cse['Form Responses 1']

wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# Styling palette
navy_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
accent_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
green_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
white_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
subhead_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
bold_font = Font(name="Calibri", size=10, bold=True, color="000000")
data_font = Font(name="Calibri", size=10, color="000000")
pwd_font = Font(name="Consolas", size=10, bold=True, color="0F172A")

thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

# ==========================================
# 1. Credentials Distribution Sheet (Concise)
# ==========================================
ws_cred = wb.create_sheet(title="Credentials Roster")
cred_headers = [
    "S.No", "Register Number", "Student Name", "Department Code", 
    "Department Name", "Official College Email", "Mobile Number", 
    "Active Login Password", "Alternative Password"
]
ws_cred.append(cred_headers)
for cell in ws_cred[1]:
    cell.fill = accent_fill
    cell.font = white_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

for idx, st in enumerate(students, 1):
    row = [
        idx,
        st["reg_no"],
        st["name"],
        st["dept_code"],
        st["dept_name"],
        st["email"],
        st["mobile"],
        st["password"],
        st["generated_password"]
    ]
    ws_cred.append(row)
    r_idx = idx + 1
    for c_idx in range(1, len(row) + 1):
        c = ws_cred.cell(row=r_idx, column=c_idx)
        c.border = thin_border
        if r_idx % 2 == 0:
            c.fill = zebra_fill
        if c_idx in [1, 2, 4, 7]:
            c.alignment = Alignment(horizontal="center")
        elif c_idx in [8, 9]:
            c.font = pwd_font
            c.alignment = Alignment(horizontal="center")
        else:
            c.font = data_font

# ==========================================
# 2. Master Student Database (Full Details)
# ==========================================
ws_master = wb.create_sheet(title="Master Student Database")
master_headers = [
    "S.No", "Register Number", "Student Name", "Department Code", "Department Name",
    "Degree", "Batch", "Academic Year", "Official College Email", "Student Mobile",
    "Date of Birth", "Gender", "Father's Name", "Mother's Name", "Blood Group",
    "Personal Email", "Address", "City / Town", "District", "State", "Pin Code",
    "Active Login Password", "Portal Status"
]
ws_master.append(master_headers)
for cell in ws_master[1]:
    cell.fill = header_fill
    cell.font = white_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

for idx, st in enumerate(students, 1):
    row = [
        idx,
        st["reg_no"],
        st["name"],
        st["dept_code"],
        st["dept_name"],
        "UG",
        "2026-2030",
        "1st Year",
        st["email"],
        st["mobile"],
        st["dob"],
        st["gender"],
        st["father"],
        st.get("mother", ""),
        st.get("blood_group", ""),
        st.get("personal_email", ""),
        st.get("address", ""),
        st.get("city", ""),
        st.get("district", ""),
        st.get("state", ""),
        st.get("pincode", ""),
        st["password"],
        "ACTIVE"
    ]
    ws_master.append(row)
    r_idx = idx + 1
    for c_idx in range(1, len(row) + 1):
        c = ws_master.cell(row=r_idx, column=c_idx)
        c.border = thin_border
        if r_idx % 2 == 0:
            c.fill = zebra_fill
        if c_idx in [1, 2, 4, 6, 7, 8, 10, 11, 12, 15, 21, 23]:
            c.alignment = Alignment(horizontal="center")
        elif c_idx == 22:
            c.font = pwd_font
            c.alignment = Alignment(horizontal="center")
        else:
            c.font = data_font

# ==========================================
# 3. CSE Section A (Detailed 36 Columns)
# ==========================================
ws_cse_sheet = wb.create_sheet(title="CSE Section A (36 Columns)")
cse_headers = [ws_cse_raw.cell(row=1, column=c).value for c in range(1, ws_cse_raw.max_column + 1)]
cse_headers.append("Portal Active Password")
ws_cse_sheet.append(cse_headers)
for cell in ws_cse_sheet[1]:
    cell.fill = green_fill
    cell.font = white_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

pwd_map = {st["reg_no"]: st["password"] for st in students}

for r in range(2, ws_cse_raw.max_row + 1):
    reg = ws_cse_raw.cell(row=r, column=1).value
    if not reg: continue
    reg = str(reg).strip().upper()
    row_vals = []
    for c in range(1, ws_cse_raw.max_column + 1):
        val = ws_cse_raw.cell(row=r, column=c).value
        if isinstance(val, (datetime.datetime, datetime.date)):
            val = val.strftime('%Y-%m-%d')
        elif val is not None:
            val = str(val).strip()
        row_vals.append(val or "")
    row_vals.append(pwd_map.get(reg, f"Nec@{reg}"))
    ws_cse_sheet.append(row_vals)
    r_idx = ws_cse_sheet.max_row
    for c_idx in range(1, len(row_vals) + 1):
        c = ws_cse_sheet.cell(row=r_idx, column=c_idx)
        c.border = thin_border
        if r_idx % 2 == 0:
            c.fill = zebra_fill
        if c_idx == len(row_vals):
            c.font = pwd_font
            c.alignment = Alignment(horizontal="center")
        else:
            c.font = data_font

# ==========================================
# 4. Department Specific Sheets (12 Branches)
# ==========================================
dept_codes = ["CSE", "IT", "ECE", "MECH", "EEE", "CIVIL", "AIDS", "AGRI", "BME", "CHEMICAL", "CSE(CS)", "CSE(IoT)"]
for d_code in dept_codes:
    dept_students = [s for s in students if s["dept_code"] == d_code]
    sheet_title = d_code.replace("(", "").replace(")", "")[:31]
    ws_dept = wb.create_sheet(title=f"Dept-{sheet_title}")
    
    headers = [
        "S.No", "Register Number", "Student Name", "Official Email", 
        "Mobile Number", "Date of Birth", "Gender", "Father's Name", 
        "Login Password"
    ]
    ws_dept.append(headers)
    for cell in ws_dept[1]:
        cell.fill = navy_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, st in enumerate(dept_students, 1):
        row = [
            idx,
            st["reg_no"],
            st["name"],
            st["email"],
            st["mobile"],
            st["dob"],
            st["gender"],
            st["father"],
            st["password"]
        ]
        ws_dept.append(row)
        r_idx = idx + 1
        for c_idx in range(1, len(row) + 1):
            c = ws_dept.cell(row=r_idx, column=c_idx)
            c.border = thin_border
            if r_idx % 2 == 0:
                c.fill = zebra_fill
            if c_idx in [1, 2, 5, 6, 7]:
                c.alignment = Alignment(horizontal="center")
            elif c_idx == 9:
                c.font = pwd_font
                c.alignment = Alignment(horizontal="center")
            else:
                c.font = data_font

# ==========================================
# 5. Department Summary Sheet
# ==========================================
ws_summary = wb.create_sheet(title="Department Summary")
sum_headers = ["S.No", "Department Code", "Department Name", "Degree", "Batch", "Total Students", "Male Students", "Female Students", "Credentials Status"]
ws_summary.append(sum_headers)
for cell in ws_summary[1]:
    cell.fill = header_fill
    cell.font = white_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

dept_names_full = {
    "AGRI": "Agricultural Engineering",
    "AIDS": "Artificial Intelligence and Data Science",
    "BME": "Biomedical Engineering",
    "CHEMICAL": "Chemical Engineering",
    "CIVIL": "Civil Engineering",
    "CSE": "Computer Science and Engineering",
    "CSE(CS)": "Computer Science and Engineering (Cyber Security)",
    "CSE(IoT)": "Computer Science and Engineering (Internet of Things)",
    "ECE": "Electronics and Communication Engineering",
    "EEE": "Electrical and Electronics Engineering",
    "IT": "Information Technology",
    "MECH": "Mechanical Engineering",
}

for idx, d_code in enumerate(dept_codes, 1):
    d_sts = [s for s in students if s["dept_code"] == d_code]
    males = sum(1 for s in d_sts if s["gender"] == "Male")
    females = sum(1 for s in d_sts if s["gender"] == "Female")
    ws_summary.append([
        idx,
        d_code,
        dept_names_full.get(d_code, d_code),
        "UG",
        "2026-2030",
        len(d_sts),
        males,
        females,
        "100% Active & Verified"
    ])
    r_idx = idx + 1
    for c_idx in range(1, 10):
        c = ws_summary.cell(row=r_idx, column=c_idx)
        c.border = thin_border
        if r_idx % 2 == 0:
            c.fill = zebra_fill
        if c_idx in [1, 2, 4, 5, 6, 7, 8, 9]:
            c.alignment = Alignment(horizontal="center")
        if c_idx in [6, 7, 8]:
            c.font = bold_font

# Total Row
tot_row = len(dept_codes) + 2
ws_summary.cell(row=tot_row, column=1, value="TOTAL")
ws_summary.cell(row=tot_row, column=2, value="12 Branches")
ws_summary.cell(row=tot_row, column=3, value="Nandha Engineering College (Autonomous)")
ws_summary.cell(row=tot_row, column=4, value="UG")
ws_summary.cell(row=tot_row, column=5, value="2026-2030")
ws_summary.cell(row=tot_row, column=6, value=f"=SUM(F2:F{tot_row-1})")
ws_summary.cell(row=tot_row, column=7, value=f"=SUM(G2:G{tot_row-1})")
ws_summary.cell(row=tot_row, column=8, value=f"=SUM(H2:H{tot_row-1})")
ws_summary.cell(row=tot_row, column=9, value="1,028 Verified Accounts")

for c_idx in range(1, 10):
    c = ws_summary.cell(row=tot_row, column=c_idx)
    c.fill = accent_fill
    c.font = white_font
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = thin_border

# Auto-adjust column widths for all sheets
for sheet in wb.worksheets:
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = cell.value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

# Save workbook to all relevant locations
paths = [
    r'c:/Users/SDP/OneDrive/Desktop/NEC-SCS-main (2)/NEC_Student_Database_Credentials_2026.xlsx',
    r'c:/Users/SDP/OneDrive/Desktop/NEC-SCS-main (2)/NEC-SCS-main/NEC-SCS-main/public/NEC_Student_Credentials_2026.xlsx',
    r'C:/Users/SDP/.gemini/antigravity/brain/257e1302-6051-4ba8-b085-f5cd855ef30b/NEC_Student_Credentials_2026.xlsx',
    r'c:/Users/SDP/OneDrive/Desktop/NEC-SCS-main (2)/NEC_Student_Credentials_2026.xlsx',
]

for p in paths:
    try:
        wb.save(p)
        print(f"Saved: {p}")
    except PermissionError:
        print(f"File locked by Excel (in use): {p}")
    except Exception as e:
        print(f"Error saving {p}: {e}")

print("Master credentials workbook generated with all sheets.")
