import openpyxl
import json
import datetime

# 1. Load users2030
wb_users = openpyxl.load_workbook('C:/Users/SDP/.gemini/antigravity/brain/257e1302-6051-4ba8-b085-f5cd855ef30b/.user_uploaded/media_1793284665788_576028e8.xlsx', data_only=True)
ws_users = wb_users['users2030']

# 2. Load CSE_A_-_DETAILS
wb_cse = openpyxl.load_workbook('C:/Users/SDP/.gemini/antigravity/brain/257e1302-6051-4ba8-b085-f5cd855ef30b/.user_uploaded/media_1793284665788_84285bee.xlsx', data_only=True)
ws_cse = wb_cse['Form Responses 1']

cse_details = {}
for r in range(2, ws_cse.max_row + 1):
    reg = ws_cse.cell(row=r, column=1).value
    if not reg: continue
    reg = str(reg).strip().upper()
    
    # Store all cols
    row_data = {}
    for c in range(1, ws_cse.max_column + 1):
        col_header = ws_cse.cell(row=1, column=c).value
        val = ws_cse.cell(row=r, column=c).value
        if isinstance(val, (datetime.datetime, datetime.date)):
            val = val.strftime('%Y-%m-%d')
        elif val is not None:
            val = str(val).strip()
        row_data[col_header] = val
    cse_details[reg] = row_data

print(f"Loaded {len(cse_details)} detailed CSE students.")

# 3. Process all 1,028 students from users2030
DEPT_NAMES = {
    "AGRI": "Agricultural Engineering",
    "AIDS": "Artificial Intelligence and Data Science",
    "AI&DS": "Artificial Intelligence and Data Science",
    "BME": "Biomedical Engineering",
    "CHEMICAL": "Chemical Engineering",
    "CIVIL": "Civil Engineering",
    "CSE": "Computer Science and Engineering",
    "CSE(CS)": "Computer Science and Engineering (Cyber Security)",
    "CSE(IoT)": "Computer Science and Engineering (Internet of Things)",
    "ECE": "Electronics and Communication Engineering",
    "EEE": "Electrical and Electronics Engineering",
    "IT": "Information Technology",
    "MECH": "Mechanical Engineering",
}

students_all = []
for r in range(2, ws_users.max_row + 1):
    reg = ws_users.cell(row=r, column=1).value
    if not reg: continue
    reg = str(reg).strip().upper()
    name = str(ws_users.cell(row=r, column=2).value or '').strip()
    email = str(ws_users.cell(row=r, column=3).value or '').strip().lower()
    raw_pwd = str(ws_users.cell(row=r, column=4).value or '').strip()
    mobile = str(ws_users.cell(row=r, column=14).value or '').strip()
    dob_val = ws_users.cell(row=r, column=15).value
    if isinstance(dob_val, (datetime.datetime, datetime.date)):
        dob = dob_val.strftime('%Y-%m-%d')
    else:
        dob = str(dob_val or '').strip()
    father = str(ws_users.cell(row=r, column=16).value or '').strip()
    gender_val = str(ws_users.cell(row=r, column=19).value or '').strip()
    dept_code = str(ws_users.cell(row=r, column=21).value or '').strip().upper()
    if dept_code == "AI&DS":
        dept_code = "AIDS"

    # Gender normalization
    if "FEMALE" in gender_val.upper():
        gender = "Female"
    elif "MALE" in gender_val.upper():
        gender = "Male"
    else:
        gender = ""

    # Password resolution:
    # If raw_pwd is valid and not NotGiven, use it; otherwise generate Nec@<RegNo>
    if raw_pwd and "notgiven" not in raw_pwd.lower():
        password = raw_pwd
    else:
        password = f"Nec@{reg}"

    # Check if CSE details exist
    cse_data = cse_details.get(reg, {})
    if cse_data:
        if cse_data.get("Date of Birth"): dob = cse_data["Date of Birth"]
        if cse_data.get("Gender"): gender = cse_data["Gender"].capitalize()
        f_fname = cse_data.get("Father's First Name") or ""
        f_lname = cse_data.get("Father's Last Name") or ""
        if f_fname or f_lname:
            father = f"{f_fname} {f_lname}".strip()

    st_obj = {
        "s_no": r - 1,
        "reg_no": reg,
        "name": name,
        "dept_code": dept_code,
        "dept_name": DEPT_NAMES.get(dept_code, dept_code),
        "email": email,
        "password": password,
        "raw_excel_password": raw_pwd,
        "generated_password": f"Nec@{reg}",
        "mobile": mobile if mobile and mobile.lower() != "notgiven" else "",
        "gender": gender,
        "dob": dob,
        "father": father,
        "mother": cse_data.get("Mother's Name", ""),
        "blood_group": cse_data.get("Blood Group", ""),
        "personal_email": cse_data.get("Email Id", ""),
        "address": cse_data.get("Address ", ""),
        "city": cse_data.get("City", ""),
        "district": cse_data.get("District", ""),
        "state": cse_data.get("State", ""),
        "pincode": cse_data.get("Pin Code", ""),
        "role": "STUDENT",
        "status": "ACTIVE",
    }
    students_all.append(st_obj)

print(f"Processed {len(students_all)} total students.")
print("Sample record 0 (AGRI):", students_all[0])
print("Sample record 290 (CS001):", [s for s in students_all if s['reg_no'] == '732226CS001'][0])

with open('public/nec_students_2026.json', 'w', encoding='utf-8') as f:
    json.dump(students_all, f, indent=2)

print("Saved updated nec_students_2026.json successfully.")

